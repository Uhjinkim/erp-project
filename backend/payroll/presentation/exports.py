from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from payroll.application.ports import EmployeeDisplay
from payroll.domain.entities import ComponentCategory, PayrollStatement
from payroll.domain.policies import payslip_rank

_THIN = Side(style="thin", color="C8CCC6")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="F4F5F1")
_NET_FILL = PatternFill("solid", fgColor="EDF5ED")
_WON = '#,##0"원"'


def payslip_filename(statement: PayrollStatement) -> str:
    period = f"{statement.period.year}{statement.period.month:02d}"
    return f"급여명세서_{period}_{statement.employee_no}.xlsx"


def build_payslip_workbook(
    statement: PayrollStatement,
    employee: EmployeeDisplay | None,
    component_names: dict[str, str],
) -> bytes:
    """Render a confirmed statement in the same layout as the payroll screen's payslip."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "급여명세서"
    for column, width in zip("ABCD", (22, 16, 22, 16), strict=True):
        sheet.column_dimensions[column].width = width

    sheet.merge_cells("A1:D1")
    sheet["A1"] = f"{statement.period.year}년 {statement.period.month}월 급여명세서"
    sheet["A1"].font = Font(size=14, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center")

    meta = [
        (
            "성명",
            employee.name if employee else "-",
            "직급",
            employee.position_name if employee else "-",
        ),
        ("사번", statement.employee_no, "지급일", statement.payment_date.isoformat()),
    ]
    for offset, row in enumerate(meta, start=3):
        for column, value in zip("ABCD", row, strict=True):
            sheet[f"{column}{offset}"] = value or "-"
        sheet[f"A{offset}"].font = Font(bold=True)
        sheet[f"C{offset}"].font = Font(bold=True)

    lines = sorted(
        (item for item in statement.items if int(item.amount) != 0),
        key=lambda item: payslip_rank(item.component_code),
    )
    earnings = [item for item in lines if item.category == ComponentCategory.EARNING]
    deductions = [item for item in lines if item.category == ComponentCategory.DEDUCTION]

    header_row = 6
    for column, label in zip(
        "ABCD", ("지급내역 (A)", "지급액", "공제내역 (B)", "공제액"), strict=True
    ):
        cell = sheet[f"{column}{header_row}"]
        cell.value = label
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL

    row = header_row
    for index in range(max(len(earnings), len(deductions), 1)):
        row = header_row + 1 + index
        for name_column, amount_column, items in (("A", "B", earnings), ("C", "D", deductions)):
            if index < len(items):
                item = items[index]
                sheet[f"{name_column}{row}"] = component_names.get(
                    item.component_code, item.component_code
                )
                sheet[f"{amount_column}{row}"] = int(item.amount)
                sheet[f"{amount_column}{row}"].number_format = _WON

    totals_row = row + 1
    sheet[f"A{totals_row}"] = "지급총액"
    sheet[f"B{totals_row}"] = int(statement.total_earnings)
    sheet[f"C{totals_row}"] = "공제액 계"
    sheet[f"D{totals_row}"] = int(statement.total_deductions)
    net_row = totals_row + 1
    sheet.merge_cells(f"A{net_row}:C{net_row}")
    sheet[f"A{net_row}"] = "차인지급액 (C)"
    sheet[f"D{net_row}"] = int(statement.net_pay)
    for target_row in (totals_row, net_row):
        for column in "ABCD":
            sheet[f"{column}{target_row}"].font = Font(bold=True)
        sheet[f"B{target_row}"].number_format = _WON
        sheet[f"D{target_row}"].number_format = _WON
    for column in "ABCD":
        sheet[f"{column}{net_row}"].fill = _NET_FILL

    for table_row in sheet.iter_rows(min_row=header_row, max_row=net_row, max_col=4):
        for cell in table_row:
            cell.border = _BORDER

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
